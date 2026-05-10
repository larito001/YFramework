"""excel-generation skill 的共享辅助。手工修改风险自负。

当前用途：
- 校验 col_type 在配表工具白名单内（与 tools/gtp/common/ConfigData.py:isValidType 保持一致）。
- 写 Excel 表头 Row1~Row6（树形表 default 在 Row7）。
- 写数据行（普通表 Row7+ / 树形表 Row8+）。
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL_NAME    = PatternFill("solid", fgColor="DDEBF7")  # Row1
HEADER_FILL_ID      = PatternFill("solid", fgColor="E2EFDA")  # Row2
HEADER_FILL_TYPE    = PatternFill("solid", fgColor="FFF2CC")  # Row3
HEADER_FILL_TARGET  = PatternFill("solid", fgColor="FCE4D6")  # Row4
HEADER_FILL_EXT     = PatternFill("solid", fgColor="F4F4F4")  # Row5
HEADER_FILL_DEFAULT = PatternFill("solid", fgColor="FFFFFF")  # Row6/7

BASE_TYPES = {"int", "uint", "float", "bool", "string"}


def is_valid_type(t: str) -> bool:
    if t in BASE_TYPES:
        return True
    parts = t.split(".")
    if len(parts) == 1:
        return False
    if parts[0] == "array":
        return is_valid_type(t[6:])
    if len(parts) == 2 and parts[0] in ("vec2", "vec3") and parts[1] in BASE_TYPES:
        return True
    if len(parts) == 3 and parts[0] == "map" and parts[1] in BASE_TYPES and parts[2] in BASE_TYPES:
        return True
    return False


def write_header(ws, columns, is_tree=False):
    """写 Row1~Row6（树形表 default 在 Row7）。"""
    bold = Font(bold=True)
    for i, c in enumerate(columns, start=1):
        cell1 = ws.cell(row=1, column=i, value=c["col_name"])
        cell1.fill = HEADER_FILL_NAME
        cell1.font = bold
        ws.cell(row=2, column=i, value=c["col_id"]).fill = HEADER_FILL_ID
        ws.cell(row=3, column=i, value=c["type"]).fill = HEADER_FILL_TYPE
        ws.cell(row=4, column=i, value=c["target"]).fill = HEADER_FILL_TARGET
        ws.cell(row=5, column=i, value=c.get("ext", "")).fill = HEADER_FILL_EXT
        default_row = 7 if is_tree else 6
        ws.cell(row=default_row, column=i, value=c.get("default", "")).fill = HEADER_FILL_DEFAULT
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = max(12, len(str(c["col_id"])) + 4)


def write_data_rows(ws, columns, rows, is_tree=False):
    start_row = 8 if is_tree else 7
    for r, row in enumerate(rows, start=start_row):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r, column=c_idx, value=val)


def build_xlsx(out_path, table_name, columns, rows, is_tree=False):
    for c in columns:
        if not is_valid_type(c["type"]):
            raise ValueError(f"Invalid type '{c['type']}' for column '{c['col_id']}'")
    if not any("key" in c["target"] for c in columns):
        raise ValueError("table must have at least one key column (target containing 'key')")
    wb = Workbook()
    ws = wb.active
    ws.title = table_name
    write_header(ws, columns, is_tree)
    write_data_rows(ws, columns, rows, is_tree)
    wb.save(out_path)
    return out_path
