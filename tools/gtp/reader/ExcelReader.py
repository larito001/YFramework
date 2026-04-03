# gtp/reader/ExcelReader.py

import os

import openpyxl

from gtp.common.ConfigData import (
    ConfigData, isValidType, getPythonValue, ARRAY_SEP
)


def _getStdType(typeStr: str) -> str:
    """标准化类型字符串"""
    typeStr = (typeStr.lower().strip()
               .replace("_", ".")
               .replace("dictionary", "map")
               .replace("vector2", "vec2")
               .replace("vector3", "vec3")
               .replace("vec3.array", "array.vec3")
               .replace("vec2.array", "array.vec2"))
    if isValidType(typeStr):
        return typeStr
    raise ValueError(f"无效的数据类型: {typeStr}")


def _getStdTarget(targetStr: str) -> str:
    """将 Row4 目标文本转为缩写。优先级: all > client > server"""
    targetStr = targetStr.lower().strip()
    if "all" in targetStr:
        return "CS"
    if "client" in targetStr:
        return "C"
    if "server" in targetStr:
        return "S"
    return ""


def readExcel(path: str) -> ConfigData | None:
    """
    读取一个 Excel 配置表，返回 ConfigData 对象。
    行结构：
      Row 1 = 列显示名, Row 2 = 列ID, Row 3 = 类型,
      Row 4 = 目标/键, Row 5 = 扩展(忽略),
      Row 6 = 默认值, Row 7+ = 数据
    树形表：Row 7 = 默认值, Row 8+ = 数据
    """
    wb = openpyxl.load_workbook(filename=path, read_only=True, data_only=True)
    st = wb.worksheets[0]
    r = ConfigData()
    r.id = os.path.splitext(os.path.basename(path))[0]
    r.name = r.id

    # ---- 1. 列检测 (Row 2) ----
    cells = st[2]
    detectionCount = 10
    noneCount = 0
    cellIndex = 1
    mapTables: dict[str, int] = {}  # {colSeq(str): actual_excel_column(int)}

    for tmpcell in cells:
        if tmpcell.data_type == 's' and tmpcell.value is not None:
            noneCount = 0
            cellVal = tmpcell.value.strip()
            if cellVal:
                r.colID.append(cellVal)
                mapTables[str(len(r.colID))] = cellIndex
            else:
                break
        else:
            noneCount += 1
            if noneCount == detectionCount:
                break
        cellIndex += 1

    colCount = len(r.colID)
    if colCount == 0:
        wb.close()
        return None

    # ---- 2. 读取列元数据 ----
    indexC: list[str] = []
    isTree = False

    for colIndex in range(1, colCount + 1):
        curCol = mapTables[str(colIndex)]

        # Row 1: 列名
        r.colName.append(str(st.cell(row=1, column=curCol).value or "").strip())

        # Row 3: 类型
        typeStr = str(st.cell(row=3, column=curCol).value or "").strip()
        r.colType.append(_getStdType(typeStr))

        # Row 4: 目标+键
        row4Str = str(st.cell(row=4, column=curCol).value or "").strip()

        extraFlag = ""
        if "main" in row4Str:
            extraFlag += "M"

        target = _getStdTarget(row4Str)
        r.colTarget.append(target + extraFlag)

        if "rowkey" in row4Str:
            r.ext.append("R")

        # 索引收集：仅客户端
        if "key" in row4Str:
            if "all" in row4Str or "client" in row4Str:
                indexC.append(r.colID[colIndex - 1])

        # 树形检测
        if "main" in row4Str or "child" in row4Str or "row" in row4Str:
            isTree = True

    if indexC:
        r.dataIndexC.append(indexC)

    # ---- 3. 读取默认值 ----
    defaultRow = 6
    for colIndex in range(1, colCount + 1):
        curCol = mapTables[str(colIndex)]
        cell = st.cell(row=defaultRow, column=curCol)

        # 树形表：Row 7 可覆盖 Row 6
        if isTree:
            cell2 = st.cell(row=7, column=curCol)
            if cell2.value is not None:
                cell = cell2

        colType = r.colType[colIndex - 1]
        cellVal = cell.value

        # 分隔符替换
        if cellVal is not None and isinstance(cellVal, str):
            if "map" in colType:
                cellVal = cellVal.replace(";", ARRAY_SEP).replace("=", "~")
            elif "vec" in colType:
                cellVal = cellVal.replace(";", "~")

        r.colDefault.append(getPythonValue(cellVal, colType, None))

    # ---- 4. 读取数据行 ----
    startRow = 7 if not isTree else 8
    rowIndex = 0
    lastRowData: list | None = None

    for row in st.rows:
        rowIndex += 1
        if len(row) == 0:
            break
        if rowIndex < startRow:
            continue

        rowData: list = []
        nullCount = 0

        for colIndex in range(1, colCount + 1):
            curCol = mapTables[str(colIndex)]
            cell = row[curCol - 1]
            cellVal = cell.value

            if cellVal is None:
                nullCount += 1

            colType = r.colType[colIndex - 1]

            # 分隔符替换
            if cellVal is not None and isinstance(cellVal, str):
                if "map" in colType:
                    cellVal = cellVal.replace(";", ARRAY_SEP).replace("=", "~")
                elif "vec" in colType:
                    cellVal = cellVal.replace(";", "~")

            # 树形继承
            if isTree and "M" in r.colTarget[colIndex - 1] and lastRowData is not None:
                rowData.append(getPythonValue(cellVal, colType, lastRowData[colIndex - 1]))
            else:
                rowData.append(getPythonValue(cellVal, colType, r.colDefault[colIndex - 1]))

        if nullCount >= colCount:
            break

        r.data.append(rowData)
        lastRowData = rowData

    wb.close()
    return r
