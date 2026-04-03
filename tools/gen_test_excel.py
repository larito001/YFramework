"""生成测试用 Excel 配置表（hero / skill / quest）"""
import os
import openpyxl

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_test", "excel", "3xlsx")
os.makedirs(OUT_DIR, exist_ok=True)

# ============================================================
# 1. hero.xlsx - 单键表，覆盖所有基础+复合类型
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "hero"

ws.append(["英雄名称", "英雄ID", "生命值", "攻击力", "速度", "稀有度", "技能列表", "属性加成", "初始位置"])
ws.append(["hero_name", "hero_id", "hp", "atk", "speed", "rarity", "skill_list", "attr_bonus", "init_pos"])
ws.append(["string", "uint", "int", "int", "float", "uint", "array.int", "map.string.int", "vec2.int"])
ws.append(["client", "client_key", "all", "all", "client", "client", "client", "client", "client"])
ws.append(["", "", "", "", "", "", "", "", ""])
ws.append(["", 0, 0, 0, 0, 0, "", "", "0~0"])
ws.append(["战神阿瑞斯", 1001, 1200, 85, 1.2, 5, "1001|1002|1003", "atk~10|def~5", "100~200"])
ws.append(["暗影刺客", 1002, 800, 120, 1.8, 4, "2001|2002", "atk~20|crit~15", "50~150"])
ws.append(["圣光牧师", 1003, 1500, 45, 1.0, 4, "3001|3002|3003", "hp~30|def~10", "200~100"])
ws.append(["烈焰法师", 1004, 900, 100, 1.1, 5, "4001|4002", "atk~25|speed~5", "150~250"])
ws.append(["铁壁骑士", 1005, 2000, 60, 0.8, 3, "5001", "def~40|hp~20", "300~100"])

heroPath = os.path.join(OUT_DIR, "hero.xlsx")
wb.save(heroPath)
print("Generated: " + heroPath)

# ============================================================
# 2. skill.xlsx - 双键表
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "skill"

ws.append(["技能名称", "英雄ID", "技能ID", "伤害", "冷却"])
ws.append(["skill_name", "hero_id", "skill_id", "damage", "cooldown"])
ws.append(["string", "uint", "uint", "int", "float"])
ws.append(["client", "client_key", "all_key", "all", "client"])
ws.append(["", "", "", "", ""])
ws.append(["", 0, 0, 0, 0])
ws.append(["火球术", 1001, 1, 100, 3.0])
ws.append(["冰冻术", 1001, 2, 80, 5.0])
ws.append(["雷击", 1002, 1, 120, 4.0])
ws.append(["治疗", 1002, 2, 0, 2.0])

skillPath = os.path.join(OUT_DIR, "skill.xlsx")
wb.save(skillPath)
print("Generated: " + skillPath)

# ============================================================
# 3. quest.xlsx - 树形表
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "quest"

ws.append(["任务名", "任务ID", "步骤ID", "描述", "奖励"])
ws.append(["quest_name", "quest_id", "step_id", "desc", "reward"])
ws.append(["string", "uint", "uint", "string", "int"])
ws.append(["client_main", "client_key_main", "all_key_child", "client_child", "client_child"])
ws.append(["", "", "", "", ""])
ws.append(["", 0, 0, "", 0])
ws.append(["默认任务", 0, 0, "默认描述", 0])
ws.append(["主线任务1", 1001, 1, "击杀10个怪物", 100])
ws.append(["", "", 2, "回城交任务", 50])
ws.append(["主线任务2", 1002, 1, "采集药草", 80])

questPath = os.path.join(OUT_DIR, "quest.xlsx")
wb.save(questPath)
print("Generated: " + questPath)

print("\nAll test excel files generated in: " + OUT_DIR)
