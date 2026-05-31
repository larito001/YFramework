# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def build(path, cols, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    hf = PatternFill('solid', fgColor='DDEBF7')
    kf = PatternFill('solid', fgColor='FCE4D6')
    bold = Font(bold=True)
    ctr = Alignment(horizontal='center')
    for c, (disp, fid, ft, tg) in enumerate(cols, 1):
        ws.cell(1, c, disp)
        ws.cell(2, c, fid).font = bold
        ws.cell(3, c, ft)
        ws.cell(4, c, tg)
        isk = 'key' in tg
        for r in range(1, 7):
            cell = ws.cell(r, c)
            cell.fill = kf if isk else hf
            cell.alignment = ctr
    for ri, row in enumerate(rows, 7):
        for ci, val in enumerate(row, 1):
            ws.cell(ri, ci, val)
    ws.freeze_panes = 'A7'
    wb.save(path)
    print('wrote', path, len(rows), 'rows')


# chest: 宝箱本体。groupSeq/rollSeq 平行数组,按开启次序(第 n 次,1-based)取第 min(n,len)-1 项;超出末尾重复最后一个
build('excel/3xlsx/chest.xlsx', [
    ('宝箱ID', 'id', 'uint', 'all key'),
    ('名称', 'name', 'string', 'all'),
    ('网格宽', 'width', 'int', 'all'),
    ('网格高', 'height', 'int', 'all'),
    ('掉落组序列', 'groupSeq', 'array.uint', 'all'),
    ('抽取次数序列', 'rollSeq', 'array.int', 'all'),
], [
    [1, '木宝箱', 8, 6, '1|2', '3|4'],
    [2, '铁宝箱', 8, 6, '2|3', '4|5'],
    [3, '黄金宝箱', 8, 6, '3', '5'],
])

# chestDrop: 掉落条目。按 groupId 分组建池;每次抽取在池内加权选 1 条,数量 [minCount,maxCount] 随机
build('excel/3xlsx/chestDrop.xlsx', [
    ('条目ID', 'id', 'uint', 'all key'),
    ('掉落组', 'groupId', 'uint', 'all'),
    ('物品ID', 'itemId', 'uint', 'all'),
    ('权重', 'weight', 'int', 'all'),
    ('最小数量', 'minCount', 'int', 'all'),
    ('最大数量', 'maxCount', 'int', 'all'),
], [
    # group1 普通
    [1, 1, 1001, 50, 1, 3], [2, 1, 3001, 50, 5, 20], [3, 1, 3002, 40, 3, 10], [4, 1, 5001, 30, 10, 100], [5, 1, 2001, 10, 1, 1],
    # group2 稀有
    [6, 2, 1002, 40, 1, 2], [7, 2, 2002, 30, 1, 1], [8, 2, 4002, 20, 1, 1], [9, 2, 5001, 50, 50, 200], [10, 2, 4001, 5, 1, 1],
    # group3 史诗
    [11, 3, 4001, 30, 1, 1], [12, 3, 2002, 40, 1, 2], [13, 3, 1002, 50, 2, 5], [14, 3, 5001, 60, 100, 500],
])